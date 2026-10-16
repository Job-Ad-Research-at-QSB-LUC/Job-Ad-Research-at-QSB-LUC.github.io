"""Extract the official full quantitative comparator and verify expert-list vintage."""
from pathlib import Path
import csv,gzip,json,hashlib
import pandas as pd
from openpyxl import load_workbook

P=Path(__file__).resolve().parent
ROOT=P.parents[2]
source=P/'sources/Related_Occupations_Research_Dataset_2025.xlsx'
target=P/'sources/onet_full_scores.csv.gz'
wb=load_workbook(source,read_only=True,data_only=True)
ws=wb['Related Occupations Data']
rows=ws.iter_rows(values_only=True);header=next(rows)
assert header[0]=='O*NET-SOC Code' and header[2]=='Related O*NET-SOC Code' and header[9]=='Relatedness Score'
expert=[];n=0;minimum=float('inf');maximum=-minimum
with gzip.open(target,'wt',newline='') as f:
    w=csv.writer(f);w.writerow(['source','target','score'])
    for r in rows:
        if not r[0]:continue
        score=float(r[9]);minimum=min(minimum,score);maximum=max(maximum,score)
        w.writerow([r[0],r[2],score]);n+=1
        if r[5] is not None:expert.append((r[0],r[2],int(r[5])))
        if n%100000==0:print('Extracted',n,flush=True)
wb.close()
ref=pd.read_csv(ROOT/'analysis/occupation_relatedness/sources/related_occupations_31_0.csv')
expected=set(ref[['O*NET-SOC Code','Related O*NET-SOC Code','Index']].itertuples(index=False,name=None))
assert set(expert)==expected and len(expert)==18460
assert n==923*922
qa={'n_pairs':n,'score_min':minimum,'score_max':maximum,'expert_list_exactly_matches_31_0':True,
    'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
    'extracted_sha256':hashlib.sha256(target.read_bytes()).hexdigest()}
(P/'sources/extraction_QA.json').write_text(json.dumps(qa,indent=2));print(qa,flush=True)
